"""
核心处理模块 - 使用 python-docx 处理 .docx 文件
插入照片到第5个表格的2x2网格，并修改日期
"""

from docx import Document
from docx.shared import Inches
import os
import re
import shutil

def process_report(template_path, photo_paths, report_date, output_path):
    """
    处理报告生成
    
    Args:
        template_path: 模版文件路径
        photo_paths: 4张照片路径列表
        report_date: 报告日期 (YYYY-MM-DD)
        output_path: 输出文件路径
    
    Returns:
        bool: 是否成功
    """
    try:
        print("[处理] 开始处理报告")
        print(f"[处理] 模版: {template_path}")
        print(f"[处理] 日期: {report_date}")
        print(f"[处理] 照片数: {len(photo_paths)}")
        
        # 1. 检查文件是否存在
        if not os.path.exists(template_path):
            print(f"[处理] ❌ 模版文件不存在: {template_path}")
            return False
        
        for i, p in enumerate(photo_paths):
            if not os.path.exists(p):
                print(f"[处理] ❌ 照片{i+1}不存在: {p}")
                return False
        
        # 2. 打开文档
        doc = Document(template_path)
        print(f"[处理] ✅ 文档已打开，共 {len(doc.tables)} 个表格")
        
        # 3. 修改日期
        new_date = format_date(report_date)
        print(f"[处理] 日期格式转换: {report_date} -> {new_date}")
        replace_date_in_doc(doc, new_date)
        
        # 4. 插入照片到表格
        if len(doc.tables) >= 5:
            target_table = doc.tables[4]  # 第5个表格 (0-indexed)
            print(f"[处理] ✅ 找到第5个表格: {len(target_table.rows)}行 x {len(target_table.columns)}列")
            insert_photos_to_table(target_table, photo_paths)
        else:
            print(f"[处理] ⚠️ 文档只有 {len(doc.tables)} 个表格，无法插入照片")
            # 仍然继续，只修改日期
        
        # 5. 保存文档
        doc.save(output_path)
        print(f"[处理] ✅ 文档已保存: {output_path}")
        
        # 6. 验证输出文件
        if os.path.exists(output_path):
            size = os.path.getsize(output_path)
            print(f"[处理] ✅ 输出文件大小: {size} 字节")
            return True
        else:
            print(f"[处理] ❌ 输出文件不存在")
            return False
        
    except Exception as e:
        print(f"[处理] ❌ 处理失败: {str(e)}")
        import traceback
        traceback.print_exc()
        return False

def format_date(date_str):
    """
    将 YYYY-MM-DD 转换为 YYYY年M月D日
    示例: 2026-06-01 -> 2026年6月1日
    """
    try:
        parts = date_str.split('-')
        if len(parts) == 3:
            year = parts[0]
            month = str(int(parts[1]))  # 去掉前导零
            day = str(int(parts[2]))    # 去掉前导零
            return f"{year}年{month}月{day}日"
        return date_str
    except:
        return date_str

def replace_date_in_doc(doc, new_date):
    """
    替换文档中的所有日期
    匹配格式: YYYY年M月D日
    """
    date_pattern = r'\d{4}年\d{1,2}月\d{1,2}日'
    count = 0
    
    # 替换段落中的日期
    for para in doc.paragraphs:
        if re.search(date_pattern, para.text):
            old_text = para.text
            para.text = re.sub(date_pattern, new_date, para.text)
            count += 1
            print(f"[日期] 段落: {old_text[:40]}... -> {para.text[:40]}...")
    
    # 替换表格中的日期
    for table_idx, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if re.search(date_pattern, para.text):
                        old_text = para.text
                        para.text = re.sub(date_pattern, new_date, para.text)
                        count += 1
                        print(f"[日期] 表格{table_idx+1}: {old_text[:30]}... -> {para.text[:30]}...")
    
    print(f"[日期] ✅ 共替换 {count} 处日期")

def insert_photos_to_table(table, photo_paths):
    """
    在表格的 2x2 网格中插入照片
    位置: (0,0), (0,1), (1,0), (1,1)
    """
    print(f"[照片] 开始插入 {len(photo_paths)} 张照片")
    
    # 检查表格尺寸
    if len(table.rows) < 2 or len(table.columns) < 2:
        print(f"[照片] ⚠️ 表格尺寸不足: {len(table.rows)}行 x {len(table.columns)}列")
        return
    
    # 定义插入位置 (2x2 网格)
    positions = [(0, 0), (0, 1), (1, 0), (1, 1)]
    
    for i, (row_idx, col_idx) in enumerate(positions):
        if i >= len(photo_paths):
            break
        
        photo_path = photo_paths[i]
        
        try:
            # 获取单元格
            cell = table.cell(row_idx, col_idx)
            
            # 清空单元格全部内容（删除所有段落和 runs）
            for para in cell.paragraphs:
                for run in para.runs:
                    run._element.getparent().remove(run._element)
            # 删除空段落占位符节点
            tc = cell._tc
            for p_elem in tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                tc.remove(p_elem)

            # 添加新段落和图片
            para = cell.add_paragraph()
            run = para.add_run()
            
            # 插入图片 (宽度 3.5 英寸，适应表格)
            run.add_picture(photo_path, width=Inches(3.5))
            
            print(f"[照片] ✅ 照片{i+1} 已插入到 ({row_idx}, {col_idx})")
            
        except Exception as e:
            print(f"[照片] ❌ 插入照片{i+1}失败: {str(e)}")
            import traceback
            traceback.print_exc()

if __name__ == '__main__':
    # 本地测试
    print("测试处理模块...")
    
    # 使用真实模版测试
    template = r'D:\微信小程序\bank_templates\上塘支行维保报告.docx'
    photos = [r'D:\微信小程序\bank-report-miniapp\test_photos\test_photo_1.jpg'] * 4  # 用同一张照片测试
    date = '2026-06-01'
    output = r'D:\微信小程序\bank-report-miniapp\test_output2.docx'
    
    if os.path.exists(template):
        success = process_report(template, photos, date, output)
        print(f"测试结果: {'成功' if success else '失败'}")
    else:
        print(f"模版文件不存在: {template}")
